"""Relatório automatizado: planilha Excel formatada + relatório executivo em HTML."""
from __future__ import annotations

from datetime import datetime

import pandas as pd
from jinja2 import Template
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import charts
from .config import REFERENCE_DATE, Paths

PCT_HINTS = ("taxa", "conv", "sinistralidade", "retencao", "pct", "participacao")
BRL_HINTS = ("premio", "sinistro_incorrido", "severidade_media", "ticket")


def build_figures(k: dict[str, pd.DataFrame], paths: Paths) -> dict[str, str]:
    f = paths.figures
    charts.premium_monthly(k["premio_mensal"], f / "premio_mensal.png")
    charts.loss_ratio(k["sinistralidade_produto"], f / "sinistralidade_produto.png")
    charts.funnel(k["funil_conversao"], f / "funil_conversao.png")
    charts.cohort_heatmap(k["retencao_coorte"], f / "retencao_coorte.png")
    charts.cancellation_by_payment(k["cancelamento_12m"], f / "cancelamento_pagamento.png")
    return {p.stem: f"figures/{p.name}" for p in f.glob("*.png")}


def write_excel(k: dict[str, pd.DataFrame], dq_summary: pd.DataFrame, paths: Paths) -> None:
    header_fill = PatternFill("solid", fgColor="2A78D6")
    with pd.ExcelWriter(paths.reports / "relatorio_kpis.xlsx", engine="openpyxl") as xw:
        sheets = {**k, "qualidade_dados": dq_summary}
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "A2"
            for j, col in enumerate(df.columns, start=1):
                cell = ws.cell(row=1, column=j)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
                fmt = "0.0%" if col.startswith(PCT_HINTS) else '"R$" #,##0.00' if col.startswith(BRL_HINTS) else None
                if fmt:
                    for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        row[0].number_format = fmt
                width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
                ws.column_dimensions[get_column_letter(j)].width = min(width, 60)


HTML = Template("""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Relatório Executivo de Seguros</title>
<style>
:root{--bg:#fcfcfb;--ink:#0b0b0b;--muted:#52514e;--line:#e4e3df;--accent:#2a78d6;--card:#ffffff}
@media (prefers-color-scheme:dark){:root{--bg:#1a1a19;--ink:#f0efec;--muted:#c3c2b7;--line:#3a3936;--card:#232321}}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}
main{max-width:1080px;margin:0 auto;padding:32px 16px}
h1{margin:0 0 4px;font-size:26px} h2{margin:40px 0 12px;font-size:19px}
.sub{color:var(--muted);margin:0 0 24px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:22px;font-weight:700;font-variant-numeric:tabular-nums}.card .l{color:var(--muted);font-size:13px}
img{max-width:100%;border:1px solid var(--line);border-radius:10px;background:#fcfcfb}
ul.ins li{margin-bottom:6px}
.tbl{overflow-x:auto} table{border-collapse:collapse;width:100%;font-size:13px;font-variant-numeric:tabular-nums}
th,td{padding:6px 10px;border-bottom:1px solid var(--line);text-align:right;white-space:nowrap}
th:first-child,td:first-child{text-align:left} th{color:var(--muted);font-weight:600}
footer{margin-top:48px;color:var(--muted);font-size:12px}
</style></head><body><main>
<h1>Relatório Executivo — Carteira de Seguros</h1>
<p class="sub">Ano de referência {{ year }} · gerado automaticamente em {{ generated }} · dados sintéticos</p>
<div class="cards">
{% for label, value in cards %}<div class="card"><div class="v">{{ value }}</div><div class="l">{{ label }}</div></div>{% endfor %}
</div>
<h2>Principais achados</h2>
<ul class="ins">{% for i in insights %}<li>{{ i }}</li>{% endfor %}</ul>
<h2>Evolução da carteira</h2><img src="{{ figs.premio_mensal }}" alt="Prêmio ganho mensal por produto">
<h2>Sinistralidade</h2><img src="{{ figs.sinistralidade_produto }}" alt="Sinistralidade por produto">
<div class="tbl">{{ tables.sinistralidade }}</div>
<h2>Jornada digital</h2><img src="{{ figs.funil_conversao }}" alt="Funil da jornada digital">
<div class="tbl">{{ tables.funil }}</div>
<h2>Retenção e cancelamento</h2><img src="{{ figs.retencao_coorte }}" alt="Retenção por coorte">
<img src="{{ figs.cancelamento_pagamento }}" alt="Cancelamento por meio de pagamento">
{% if figs.modelo_lift_decil %}<h2>Modelo de risco de cancelamento</h2>
<p>{{ model_text }}</p><img src="{{ figs.modelo_lift_decil }}" alt="Lift por decil">{% endif %}
<h2>Qualidade de dados</h2>
<p>{{ dq_text }}</p><div class="tbl">{{ tables.dq }}</div>
<footer>Fonte: data mart <code>data/warehouse.db</code> · consultas em <code>src/insurance_analytics/sql/kpis/</code></footer>
</main></body></html>""")


def _fmt_brl(v: float) -> str:
    return f"R$ {v:,.0f}".replace(",", ".")


def _fmt_table(df: pd.DataFrame) -> str:
    d = df.copy()
    for c in d.columns:
        if c.startswith(PCT_HINTS):
            d[c] = d[c].map(lambda v: f"{v:.1%}".replace(".", ",") if pd.notna(v) else "–")
        elif c.startswith(BRL_HINTS):
            d[c] = d[c].map(_fmt_brl)
    return d.to_html(index=False, border=0)


def insights(k: dict[str, pd.DataFrame]) -> list[str]:
    """Achados gerados a partir dos números, para o texto acompanhar os dados a cada execução."""
    lr = k["sinistralidade_produto"]
    worst, best = lr.iloc[0], lr.iloc[-1]
    canc = k["cancelamento_12m"].groupby("meio_pagamento")[["canceladas_12m", "apolices"]].sum()
    canc_rate = canc["canceladas_12m"] / canc["apolices"]
    fun = k["funil_conversao"]
    steps = {
        "oferta → simulação": fun["simulacoes"].sum() / fun["ofertas_visualizadas"].sum(),
        "simulação → proposta": fun["propostas"].sum() / fun["simulacoes"].sum(),
        "proposta → contratação": fun["contratacoes"].sum() / fun["propostas"].sum(),
    }
    bottleneck = min(steps, key=steps.get)
    cross = k["cross_sell"]
    mono = cross.loc[cross["qtd_produtos"] == 1, "pct_clientes"].iloc[0]
    ret = k["retencao_coorte"].query("meses_desde_inicio == 12")["retencao"].mean()
    pct = lambda v: f"{v:.0%}"  # noqa: E731
    return [
        f"{worst['produto']} tem a maior sinistralidade ({pct(worst['sinistralidade'])}), contra "
        f"{pct(best['sinistralidade'])} em {best['produto']}: revisar precificação/aceitação de {worst['produto']}.",
        f"Pagamento via Pix cancela {pct(canc_rate['pix'])} em 12 meses, contra {pct(canc_rate['debito_conta'])} "
        f"no débito em conta: incentivar débito automático na contratação.",
        f"O maior gargalo do funil é {bottleneck} ({pct(steps[bottleneck])} de passagem).",
        f"{pct(mono)} dos clientes ativos têm um único produto: base relevante para cross-sell.",
        f"Em média, {pct(ret)} das apólices seguem ativas 12 meses após a contratação.",
    ]


def write_html(k: dict[str, pd.DataFrame], dq_summary: pd.DataFrame, figs: dict[str, str], paths: Paths,
               model_metrics: dict | None = None) -> None:
    r = k["resumo_executivo"].iloc[0]
    cards = [
        ("Apólices ativas", f"{int(r['apolices_ativas']):,}".replace(",", ".")),
        ("Clientes ativos", f"{int(r['clientes_ativos']):,}".replace(",", ".")),
        ("Prêmio ganho no ano", _fmt_brl(r["premio_ganho"])),
        ("Sinistralidade", f"{r['sinistralidade']:.1%}".replace(".", ",")),
        ("Ticket médio mensal", f"R$ {r['ticket_medio']:.2f}".replace(".", ",")),
        ("Produtos por cliente", f"{r['produtos_por_cliente']:.2f}".replace(".", ",")),
    ]
    failed = dq_summary[dq_summary["status"] == "FALHA"]
    dq_text = (f"{len(failed)} de {len(dq_summary)} regras encontraram problemas na base bruta "
               f"({int(failed['linhas_com_falha'].sum())} linhas). Linhas com falha bloqueante foram para quarentena; "
               "as demais foram corrigidas no ETL.")
    model_text = ""
    if model_metrics:
        m = model_metrics
        model_text = (f"Regressão logística com AUC de {m['auc_teste']:.3f} na base de teste. Os 10% de apólices com "
                      f"maior score concentram {m['captura_top_decil']:.0%} dos cancelamentos "
                      f"(lift de {m['lift_top_decil']:.1f}x), permitindo ação de retenção focada.")
    html = HTML.render(
        year=REFERENCE_DATE.year, generated=datetime.now().strftime("%d/%m/%Y %H:%M"), cards=cards,
        insights=insights(k), figs=figs, dq_text=dq_text, model_text=model_text,
        tables={
            "sinistralidade": _fmt_table(k["sinistralidade_produto"]),
            "funil": _fmt_table(k["funil_conversao"]),
            "dq": _fmt_table(dq_summary[["regra", "dimensao", "severidade", "linhas_com_falha", "pct_falha", "status"]]),
        },
    )
    (paths.reports / "relatorio_executivo.html").write_text(html, encoding="utf-8")
